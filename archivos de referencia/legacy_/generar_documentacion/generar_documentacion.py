import os
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor

def generar_excel_contexto():
    print("📊 Generando Excel de Contexto y Arquitectura...")
    
    # 1. Pestaña: Arquitectura y Estructura
    data_arquitectura = [
        {"Componente": "Patrón de Diseño", "Descripción": "Arquitectura Lambda modificada. Consta de una Capa Speed (Tiempo Real / Tablas Vivas) y una Capa Batch (Histórico / Cierre)."},
        {"Componente": "Extracción (RPA)", "Descripción": "Playwright + Python simulando interacción humana con el CRM Cari AI para extraer reportes CSV por franjas horarias."},
        {"Componente": "Orquestador", "Descripción": "Script bucle_maestro.py que controla el reloj simulado, lanza el RPA, el ETL, las Inteligencias Artificiales y los cierres escalonados (15:00, 16:00, 17:00)."},
        {"Componente": "Procesamiento (ETL)", "Descripción": "Pandas (motor_etl.py) limpia, cruza y estandariza la data cruda, subiéndola a la base de datos viva."},
        {"Componente": "Base de Datos", "Descripción": "PostgreSQL alojado en Supabase. Tablas vivas para RT y tablas históricas para EOD."},
        {"Componente": "Machine Learning", "Descripción": "Modelos No Supervisados (Stateless). Isolation Forest para detección de anomalías y K-Means para perfilamiento de asesores. Entrenados con Sliding Window de 30 días."},
        {"Componente": "Frontend / Dashboard", "Descripción": "FastAPI (app.py) consumiendo datos de Supabase. HTML/Bootstrap/Jinja2 (templates) para visualización web."}
    ]
    
    # 2. Pestaña: Flujo Operativo (End-to-End)
    data_flujo = [
        {"Paso": 1, "Fase": "RPA", "Acción": "bucle_maestro.py avanza el reloj simulado y lanza bot_scraper.py para extraer los reportes de la última hora desde Cari AI."},
        {"Paso": 2, "Fase": "ETL", "Acción": "motor_etl.py lee los CSV descargados, limpia formatos, cruza datos y hace INSERT en las tablas vivas (Capa Speed) de Supabase."},
        {"Paso": 3, "Fase": "ML (Horario)", "Acción": "Se ejecuta Isolation Forest sobre la data de la última hora para detectar anomalías de operación (micro) y de tráfico (macro)."},
        {"Paso": 4, "Fase": "Consumo (RT)", "Acción": "Los tableros web de Tiempo Real consultan las tablas vivas y muestran el avance diario de las campañas."},
        {"Paso": 5, "Fase": "Cierres Escalonados", "Acción": "A las 15:00, 16:00 y 17:00, la data viva se copia (INSERT INTO ... SELECT) a las tablas históricas correspondientes."},
        {"Paso": 6, "Fase": "ML (Cierre)", "Acción": "Inmediatamente tras cada cierre, se ejecuta K-Means para perfilar a los asesores basándose en los datos finales."},
        {"Paso": 7, "Fase": "Reinicio Nocturno", "Acción": "Tras 5 minutos de pausa al final del día, se ejecuta un TRUNCATE a las tablas vivas y se borra la fecha de prueba del histórico para dejar el sistema listo para el día siguiente."}
    ]
    
    # 3. Pestaña: Pase a Producción (CORREGIDO)
    data_produccion = [
        {"Archivo": "bucle_maestro.py", "Línea/Variable": "TIEMPO_ESPERA_ENTRE_HORAS", "Cambio a realizar": "Descomentar `TIEMPO_ESPERA_ENTRE_HORAS = 3600` para que el script espere 1 hora real entre ejecuciones, eliminando los 5 segundos de simulación."},
        {"Archivo": "bucle_maestro.py", "Línea/Variable": "Lógica Reposo Nocturno", "Cambio a realizar": "Descomentar la lógica de reposo nocturno que calcula los segundos faltantes hasta las 23:59:59 (reemplazando el sleep de 300 segundos)."},
        {"Archivo": "bucle_maestro.py", "Línea/Variable": "reiniciar_entorno()", "Cambio a realizar": "En la función `reiniciar_entorno`, ELIMINAR O COMENTAR todos los comandos `DELETE FROM` (La Lógica de Prueba). En producción NUNCA se borra el histórico."},
        {"Archivo": "app.py y app_dashboard.py", "Línea/Variable": "fecha_hoy_real", "Cambio a realizar": "Reemplazar `fecha_hoy_real = '2025-06-26'` por la variable de tiempo real `datetime.now(tz).strftime('%Y-%m-%d')` para que los tableros lean el día actual."},
        {"Archivo": "Archivos de IA (*_ml.py)", "Línea/Variable": "Consultas SQL", "Cambio a realizar": "Modificar los comandos `SELECT * FROM hist_...` agregando `WHERE date >= CURRENT_DATE - INTERVAL '30 days'` para aplicar la ventana deslizante."}
    ]

    with pd.ExcelWriter('Contexto_Arquitectura_Apex.xlsx', engine='openpyxl') as writer:
        pd.DataFrame(data_arquitectura).to_excel(writer, sheet_name='Arquitectura_Estructura', index=False)
        pd.DataFrame(data_flujo).to_excel(writer, sheet_name='Flujo_Operativo', index=False)
        pd.DataFrame(data_produccion).to_excel(writer, sheet_name='Pase_a_Produccion', index=False)
    print("   ✅ Excel generado con éxito: Contexto_Arquitectura_Apex.xlsx")


def generar_word_codigo():
    print("📝 Generando Word con el código fuente (excluyendo legacy_ y carpetas de sistema)...")
    
    doc = Document()
    
    # Configuración de estilos
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)
    
    doc.add_heading('Código Fuente - Trabajo de Grado', 0)
    
    # Carpetas permitidas para escanear
    carpetas_permitidas = ['.', 'templates', 'templates_ia', 'static']
    # Carpetas a excluir explícitamente
    carpetas_ignoradas = ['legacy_', '__pycache__', 'venv', '.git', 'base_datos_interna', 'descargas_bot', 'historico_procesados']
    # Extensiones de archivos a documentar
    extensiones_validas = ['.py', '.html', '.css', '.js', '.txt']
    
    archivos_procesados = 0

    for root, dirs, files in os.walk('.'):
        # Modificamos dirs in-place para que os.walk ignore las carpetas que no queremos
        dirs[:] = [d for d in dirs if d not in carpetas_ignoradas]
        
        # Filtrar solo la raíz y las carpetas permitidas
        carpeta_actual = os.path.basename(root)
        if root != '.' and carpeta_actual not in carpetas_permitidas:
            continue
            
        for file in files:
            if any(file.endswith(ext) for ext in extensiones_validas):
                file_path = os.path.join(root, file)
                
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        codigo = f.read()
                        
                    doc.add_heading(f"Archivo: {file_path}", level=1)
                    
                    # Formato de código (Courier New, tamaño 8, color gris oscuro)
                    p = doc.add_paragraph()
                    run = p.add_run(codigo)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor(50, 50, 50)
                    
                    archivos_procesados += 1
                except Exception as e:
                    print(f"   ⚠️ No se pudo leer el archivo {file_path}: {e}")

    doc.save('Codigo_Fuente_Proyecto.docx')
    print(f"   ✅ Word generado con éxito: Codigo_Fuente_Proyecto.docx ({archivos_procesados} archivos documentados)")

if __name__ == "__main__":
    print("🚀 INICIANDO GENERADOR DE DOCUMENTACIÓN...")
    generar_excel_contexto()
    generar_word_codigo()
    print("🎉 ¡Proceso finalizado! Los archivos están listos en tu carpeta.")