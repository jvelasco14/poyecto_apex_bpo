import os
from docx import Document
from docx.shared import Pt, RGBColor

def generar_word_legacy():
    print("📝 Generando Word con el código de la carpeta 'legacy_'...")
    
    doc = Document()
    
    # Configuración de estilos generales
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)
    
    doc.add_heading('Anexo: Scripts de Carga y Configuración (Legacy)', 0)
    
    # Definir explícitamente la carpeta objetivo
    carpeta_objetivo = 'legacy_'
    
    # Extensiones de archivos a documentar (incluimos .sql para creación de tablas)
    extensiones_validas = ['.py', '.sql', '.txt', '.md', '.json']
    
    archivos_procesados = 0

    # Verificamos si la carpeta existe en el directorio actual
    if not os.path.exists(carpeta_objetivo):
        print(f"❌ Error: No se encontró la carpeta '{carpeta_objetivo}'. Asegúrate de estar en la raíz del proyecto.")
        return

    # Recorrer únicamente la carpeta legacy_
    for root, dirs, files in os.walk(carpeta_objetivo):
        for file in files:
            if any(file.endswith(ext) for ext in extensiones_validas):
                file_path = os.path.join(root, file)
                
                try:
                    # Leemos el contenido del archivo
                    with open(file_path, 'r', encoding='utf-8') as f:
                        codigo = f.read()
                        
                    # Añadimos el título del archivo
                    doc.add_heading(f"Archivo: {file_path}", level=1)
                    
                    # Formato de código (Courier New, tamaño 8, color gris oscuro)
                    p = doc.add_paragraph()
                    run = p.add_run(codigo)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor(50, 50, 50)
                    
                    archivos_procesados += 1
                except Exception as e:
                    print(f"   ⚠️ No se pudo leer el archivo {file_path}: {e}")

    # Guardar el documento con un nombre diferente para no sobreescribir el anterior
    nombre_archivo = 'Codigo_Legacy_Proyecto.docx'
    doc.save(nombre_archivo)
    print(f"   ✅ Word generado con éxito: {nombre_archivo} ({archivos_procesados} archivos documentados)")

if __name__ == "__main__":
    print("🚀 INICIANDO ESCANEO DE CARPETA LEGACY...")
    generar_word_legacy()
    print("🎉 ¡Proceso finalizado!")