import os
from sqlalchemy import create_engine, text
from docx import Document
from docx.shared import Pt, RGBColor

# Configuración de conexión a Supabase (tomada de tu bucle_maestro.py)
DB_URL = "postgresql://postgres.vdvmyueufdfviztmribf:Velasco14062000**@aws-1-us-east-1.pooler.supabase.com:5432/postgres"

def generar_word_bd():
    print("🗄️ Conectando a Supabase para extraer estructura de la Base de Datos...")
    
    try:
        engine = create_engine(DB_URL)
        doc = Document()
        
        # Estilos generales del documento
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(10)
        
        doc.add_heading('Diccionario de Datos y Arquitectura Supabase - Apex BPO', 0)
        doc.add_paragraph("Este documento ha sido autogenerado conectándose directamente al motor PostgreSQL de Supabase. Contiene el esquema de tablas, definiciones de vistas y configuración de autenticación.")
        
        with engine.connect() as conn:
            
            # ==========================================
            # 1. TABLAS Y COLUMNAS
            # ==========================================
            doc.add_heading('1. Tablas en la Base de Datos (Esquema Public)', level=1)
            query_tablas = text("""
                SELECT table_name 
                FROM information_schema.tables 
                WHERE table_schema = 'public' AND table_type = 'BASE TABLE'
                ORDER BY table_name;
            """)
            tablas = conn.execute(query_tablas).fetchall()
            
            for (tabla,) in tablas:
                doc.add_heading(f"Tabla: {tabla}", level=2)
                
                query_columnas = text(f"""
                    SELECT column_name, data_type, character_maximum_length 
                    FROM information_schema.columns 
                    WHERE table_schema = 'public' AND table_name = '{tabla}'
                    ORDER BY ordinal_position;
                """)
                columnas = conn.execute(query_columnas).fetchall()
                
                # Crear tabla en Word
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Nombre de Columna'
                hdr_cells[1].text = 'Tipo de Dato'
                hdr_cells[2].text = 'Longitud Máx'
                
                # Dar formato a las celdas de encabezado
                for cell in hdr_cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                
                for col in columnas:
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(col[0])
                    row_cells[1].text = str(col[1])
                    row_cells[2].text = str(col[2]) if col[2] else 'N/A'
                    
                doc.add_paragraph("")  # Salto de línea
            
            # ==========================================
            # 2. VISTAS SQL (VIEWS)
            # ==========================================
            doc.add_heading('2. Vistas (Views) de Tiempo Real y EOD', level=1)
            doc.add_paragraph("Las siguientes vistas combinan datos de la Capa Viva y la Capa Histórica (Arquitectura Lambda):")
            
            query_vistas = text("""
                SELECT table_name, view_definition 
                FROM information_schema.views 
                WHERE table_schema = 'public'
                ORDER BY table_name;
            """)
            vistas = conn.execute(query_vistas).fetchall()
            
            for vista_nombre, vista_def in vistas:
                doc.add_heading(f"Vista: {vista_nombre}", level=2)
                p = doc.add_paragraph()
                run = p.add_run(str(vista_def).strip())
                run.font.name = 'Courier New'
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(50, 50, 50)
            
            # ==========================================
            # 3. AUTENTICACIÓN Y ROLES (SUPABASE AUTH)
            # ==========================================
            doc.add_heading('3. Módulo de Autenticación y Usuarios', level=1)
            doc.add_paragraph("Registros del módulo de System Administrator de la plataforma:")
            
            try:
                # 3.1 Usuarios en la plataforma visual (Tabla pública)
                doc.add_heading('Usuarios Frontend (usuarios_plataforma)', level=2)
                query_usuarios = text("SELECT id, nombre, email, rol, estado FROM usuarios_plataforma ORDER BY id;")
                usuarios = conn.execute(query_usuarios).fetchall()
                
                t_users = doc.add_table(rows=1, cols=5)
                t_users.style = 'Table Grid'
                h_users = t_users.rows[0].cells
                h_users[0].text = 'ID'
                h_users[1].text = 'Nombre'
                h_users[2].text = 'Email'
                h_users[3].text = 'Rol'
                h_users[4].text = 'Estado'
                
                for u in usuarios:
                    r_users = t_users.add_row().cells
                    r_users[0].text = str(u[0])
                    r_users[1].text = str(u[1])
                    r_users[2].text = str(u[2])
                    r_users[3].text = str(u[3])
                    r_users[4].text = str(u[4])
                    
            except Exception as e:
                doc.add_paragraph(f"⚠️ No se pudo cargar 'usuarios_plataforma': {e}")
                
            try:
                # 3.2 Autenticación Backend de Supabase (auth.users)
                doc.add_heading('Perfiles en Supabase Auth (Backend Security)', level=2)
                doc.add_paragraph("Esta tabla es gestionada automáticamente por Supabase para la criptografía de contraseñas y JWT tokens:")
                
                query_auth = text("SELECT id, email, created_at FROM auth.users ORDER BY created_at;")
                auth_users = conn.execute(query_auth).fetchall()
                
                t_auth = doc.add_table(rows=1, cols=3)
                t_auth.style = 'Table Grid'
                h_auth = t_auth.rows[0].cells
                h_auth[0].text = 'ID UUID (Supabase)'
                h_auth[1].text = 'Email Autenticado'
                h_auth[2].text = 'Fecha de Creación'
                
                for au in auth_users:
                    r_auth = t_auth.add_row().cells
                    r_auth[0].text = str(au[0])
                    r_auth[1].text = str(au[1])
                    r_auth[2].text = str(au[2])[:19] # Formatear la fecha
                    
            except Exception as e:
                doc.add_paragraph("⚠️ Nota: Las credenciales actuales no permiten leer la tabla interna 'auth.users', el control se lleva mediante 'usuarios_plataforma'.")

        # Guardar el documento
        nombre_archivo = 'Diccionario_Base_Datos_Apex.docx'
        doc.save(nombre_archivo)
        print(f"   ✅ ¡Misión cumplida! Diccionario generado con éxito: {nombre_archivo}")
        
    except Exception as e:
        print(f"❌ Error crítico al conectar a la base de datos: {e}")

if __name__ == "__main__":
    generar_word_bd()